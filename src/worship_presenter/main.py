import os
import sys
import tempfile
from dataclasses import replace

from pptx import Presentation as PptxPresentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from PySide6.QtCore import Qt
from PySide6.QtGui import QColor, QFont, QImage, QPixmap
from PySide6.QtWidgets import (
    QApplication,
    QCheckBox,
    QColorDialog,
    QComboBox,
    QFileDialog,
    QFontComboBox,
    QFormLayout,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QListWidget,
    QMainWindow,
    QPushButton,
    QScrollArea,
    QSlider,
    QSpinBox,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

from models import Background, GlobalDefaults, Slide, SlideType, Style
from renderer import SlideRenderer, apply_brightness_contrast, apply_text_case
from presentation import Presentation, SequenceResolver
from sequence_editor import SequenceEditor
from project_io import save_project, load_project

from presenter_defaults import (
    PresenterDefaultsWindow,
    load_presenter_settings,
)


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()

        self.setWindowTitle("Worship Presenter")
        self.resize(1400, 850)

        self.slides = []
        self.current_slide = None
        self.renderer = SlideRenderer()
        self.global_defaults = GlobalDefaults()
        self.current_project_path = None

        self.presentation = Presentation()

        # Load saved presenter defaults
        load_presenter_settings(
            self.global_defaults,
            self.presentation
        )

        self.sequence_window = None

        # Values chosen in dialogs (color picker, file picker) that haven't
        # been committed to a slide/default yet -- committed on Save.
        self._pending_primary_color = self.global_defaults.primary_style.font_color
        self._pending_secondary_color = self.global_defaults.secondary_style.font_color
        self._pending_background_path = ""

        # Cache of the original (unadjusted) background image, downscaled
        # for a fast live preview.
        self._cached_original_path = None
        self._cached_thumbnail_original = None

        self.build_ui()

        # Reflect the starting global defaults in the editor.
        self._load_primary_style(self.global_defaults.primary_style)
        self._load_secondary_style(self.global_defaults.secondary_style)
        self._apply_background_to_widgets(self.global_defaults.background)

    # =========================================================
    # MAIN UI
    # =========================================================

    def build_ui(self):
        # Top Menu Bar
        menu_bar = self.menuBar()

        # File Menu
        file_menu = menu_bar.addMenu("File")

        open_action = file_menu.addAction("Open Project")
        save_action = file_menu.addAction("Save Project")
        save_action.setShortcut("Ctrl+S")

        file_menu.addSeparator()

        export_action = file_menu.addAction("Export PowerPoint")
        export_action.triggered.connect(self.export_powerpoint)

        open_action.triggered.connect(self.open_project_dialog)
        save_action.triggered.connect(self.save_project_dialog)

        # Presenter Defaults Menu
        defaults_action = menu_bar.addAction("Presenter Defaults")
        defaults_action.triggered.connect(self.open_presenter_defaults)
        
        central = QWidget()
        self.setCentralWidget(central)

        main_layout = QHBoxLayout(central)

        # -----------------------------------------------------
        # LEFT SIDEBAR
        # -----------------------------------------------------

        sidebar = QWidget()
        sidebar.setFixedWidth(220)

        sidebar_layout = QVBoxLayout(sidebar)

        title = QLabel("Slides")
        title.setStyleSheet(
            "font-size: 20px; font-weight: bold;"
        )

        self.add_button = QPushButton("+ Add Slide")
        self.delete_button = QPushButton("Delete Slide")

        self.sequence_button = QPushButton(
            "Presentation Sequence"
        )

        self.slide_list = QListWidget()

        sidebar_layout.addWidget(title)
        sidebar_layout.addWidget(self.add_button)
        sidebar_layout.addWidget(self.delete_button)
        sidebar_layout.addWidget(self.sequence_button)
        sidebar_layout.addWidget(self.slide_list)

        self.add_button.clicked.connect(self.add_slide)
        self.delete_button.clicked.connect(self.delete_slide)
        self.slide_list.currentRowChanged.connect(
            self.select_slide
        )

        # -----------------------------------------------------
        # EDITOR
        # -----------------------------------------------------

        editor_container = QWidget()
        editor_layout = QVBoxLayout(editor_container)

        editor_title = QLabel("Slide Editor")
        editor_title.setStyleSheet(
            "font-size: 20px; font-weight: bold;"
        )

        editor_layout.addWidget(editor_title)

        # Slide type
        type_layout = QHBoxLayout()

        type_label = QLabel("Slide Type:")

        self.type_combo = QComboBox()

        for slide_type in SlideType:
            self.type_combo.addItem(
                slide_type.value.title(),
                slide_type,
            )

        type_layout.addWidget(type_label)
        type_layout.addWidget(self.type_combo)
        type_layout.addStretch()

        editor_layout.addLayout(type_layout)

        # -----------------------------------------------------
        # STYLE: CUSTOM TOGGLE + SAVE-AS-DEFAULT
        # -----------------------------------------------------

        style_header_row = QHBoxLayout()

        self.custom_style_checkbox = QCheckBox(
            "Use custom text formatting for this slide"
        )

        self.set_style_default_button = QPushButton(
            "Save as Default Formatting (All Slides)"
        )

        style_header_row.addWidget(self.custom_style_checkbox)
        style_header_row.addStretch()
        style_header_row.addWidget(self.set_style_default_button)

        editor_layout.addLayout(style_header_row)

        # -----------------------------------------------------
        # LANGUAGE EDITORS
        # -----------------------------------------------------

        language_layout = QHBoxLayout()

        # Primary
        primary_group = QGroupBox("Primary Language")
        primary_layout = QVBoxLayout(primary_group)

        self.primary_text = QTextEdit()
        self.primary_text.setPlaceholderText(
            "Paste primary language lyrics here..."
        )

        primary_layout.addWidget(self.primary_text)

        self.primary_font = QFontComboBox()

        self.primary_size = QSpinBox()
        self.primary_size.setRange(8, 150)
        self.primary_size.setValue(36)

        self.primary_color_button = QPushButton(
            "Choose Text Color"
        )

        self.primary_bold = QCheckBox("Bold")
        self.primary_italic = QCheckBox("Italic")

        self.primary_case = QComboBox()

        self.primary_case.addItem("Normal", "normal")
        self.primary_case.addItem("UPPERCASE", "upper")
        self.primary_case.addItem("lowercase", "lower")
        self.primary_case.addItem("Sentence case", "sentence")
        self.primary_case.addItem("Title Case", "title")

        primary_style = QFormLayout()

        primary_style.addRow(
            "Font:",
            self.primary_font,
        )

        primary_style.addRow(
            "Text Case:",
            self.primary_case
        )

        primary_style.addRow(
            "Size:",
            self.primary_size,
        )

        primary_style.addRow(
            self.primary_color_button
        )

        primary_style.addRow(
            self.primary_bold
        )

        primary_style.addRow(
            self.primary_italic
        )

        primary_layout.addLayout(primary_style)

        # Secondary
        secondary_group = QGroupBox("Secondary Language")
        secondary_layout = QVBoxLayout(secondary_group)

        self.secondary_text = QTextEdit()
        self.secondary_text.setPlaceholderText(
            "Paste secondary language lyrics here..."
        )

        secondary_layout.addWidget(self.secondary_text)

        self.secondary_font = QFontComboBox()

        self.secondary_size = QSpinBox()
        self.secondary_size.setRange(8, 150)
        self.secondary_size.setValue(30)

        self.secondary_color_button = QPushButton(
            "Choose Text Color"
        )

        self.secondary_bold = QCheckBox("Bold")
        self.secondary_italic = QCheckBox("Italic")

        self.secondary_case = QComboBox()

        self.secondary_case.addItem("Normal", "normal")
        self.secondary_case.addItem("UPPERCASE", "upper")
        self.secondary_case.addItem("lowercase", "lower")
        self.secondary_case.addItem("Sentence case", "sentence")
        self.secondary_case.addItem("Title Case", "title")

        secondary_style = QFormLayout()

        secondary_style.addRow(
            "Font:",
            self.secondary_font,
        )

        secondary_style.addRow(
            "Text Case:",
            self.secondary_case
        )

        secondary_style.addRow(
            "Size:",
            self.secondary_size,
        )

        secondary_style.addRow(
            self.secondary_color_button
        )

        secondary_style.addRow(
            self.secondary_bold
        )

        secondary_style.addRow(
            self.secondary_italic
        )

        secondary_layout.addLayout(secondary_style)

        language_layout.addWidget(primary_group)
        language_layout.addWidget(secondary_group)

        editor_layout.addLayout(language_layout, 1)

        # -----------------------------------------------------
        # BACKGROUND
        # -----------------------------------------------------

        background_group = QGroupBox("Background")
        background_outer = QVBoxLayout(background_group)

        # Custom toggle + save-as-default
        bg_header_row = QHBoxLayout()

        self.custom_background_checkbox = QCheckBox(
            "Use custom background for this slide"
        )

        self.set_background_default_button = QPushButton(
            "Save as Default Background (All Slides)"
        )

        bg_header_row.addWidget(self.custom_background_checkbox)
        bg_header_row.addStretch()
        bg_header_row.addWidget(self.set_background_default_button)

        background_outer.addLayout(bg_header_row)

        # Image chooser + filename + live preview thumbnail
        bg_image_row = QHBoxLayout()

        self.background_button = QPushButton(
            "Choose Image"
        )

        self.background_path_label = QLabel("No image selected")

        self.background_preview = QLabel()
        self.background_preview.setFixedSize(200, 120)
        self.background_preview.setStyleSheet(
            "border: 1px solid #888888; background: #222222; color: #aaaaaa;"
        )
        self.background_preview.setAlignment(Qt.AlignCenter)
        self.background_preview.setScaledContents(True)
        self.background_preview.setText("No Image")

        bg_image_row.addWidget(self.background_button)
        bg_image_row.addWidget(self.background_path_label, 1)
        bg_image_row.addWidget(self.background_preview)

        background_outer.addLayout(bg_image_row)

        # Brightness row: slider + numeric spin box
        brightness_row = QHBoxLayout()

        self.brightness_slider = QSlider(Qt.Horizontal)
        self.brightness_slider.setRange(-100, 100)
        self.brightness_slider.setValue(0)

        self.brightness_spin = QSpinBox()
        self.brightness_spin.setRange(-100, 100)
        self.brightness_spin.setValue(0)

        brightness_row.addWidget(QLabel("Brightness"))
        brightness_row.addWidget(self.brightness_slider, 1)
        brightness_row.addWidget(self.brightness_spin)

        background_outer.addLayout(brightness_row)

        # Contrast row: slider + numeric spin box
        contrast_row = QHBoxLayout()

        self.contrast_slider = QSlider(Qt.Horizontal)
        self.contrast_slider.setRange(0, 200)
        self.contrast_slider.setValue(100)

        self.contrast_spin = QSpinBox()
        self.contrast_spin.setRange(0, 200)
        self.contrast_spin.setValue(100)

        contrast_row.addWidget(QLabel("Contrast"))
        contrast_row.addWidget(self.contrast_slider, 1)
        contrast_row.addWidget(self.contrast_spin)

        background_outer.addLayout(contrast_row)

        editor_layout.addWidget(background_group)

        # -----------------------------------------------------
        # SAVE / PREVIEW
        # -----------------------------------------------------

        bottom_layout = QHBoxLayout()

        self.save_button = QPushButton("Apply Changes")
        self.preview_button = QPushButton("Preview")

        bottom_layout.addWidget(self.save_button)
        bottom_layout.addWidget(self.preview_button)
        bottom_layout.addStretch()

        editor_layout.addLayout(bottom_layout)

        # -----------------------------------------------------
        # ADD TO MAIN WINDOW
        # -----------------------------------------------------

        main_layout.addWidget(sidebar)
        main_layout.addWidget(editor_container, 1)

        # -----------------------------------------------------
        # SIGNALS
        # -----------------------------------------------------

        self.type_combo.currentIndexChanged.connect(
            self.update_slide_type
        )

        self.save_button.clicked.connect(
            self.save_current_slide
        )

        self.preview_button.clicked.connect(
            self.preview_slide
        )

        self.primary_color_button.clicked.connect(
            lambda: self.choose_color("primary")
        )

        self.secondary_color_button.clicked.connect(
            lambda: self.choose_color("secondary")
        )

        self.background_button.clicked.connect(
            self.choose_background
        )

        self.custom_style_checkbox.toggled.connect(
            self.on_custom_style_toggled
        )

        self.custom_background_checkbox.toggled.connect(
            self.on_custom_background_toggled
        )

        self.set_style_default_button.clicked.connect(
            self.set_style_as_default
        )

        self.set_background_default_button.clicked.connect(
            self.set_background_as_default
        )

        # Keep slider <-> spin box in sync both ways.
        self.brightness_slider.valueChanged.connect(
            self.brightness_spin.setValue
        )
        self.brightness_spin.valueChanged.connect(
            self.brightness_slider.setValue
        )

        self.contrast_slider.valueChanged.connect(
            self.contrast_spin.setValue
        )
        self.contrast_spin.valueChanged.connect(
            self.contrast_slider.setValue
        )

        # Live preview updates whenever brightness/contrast changes.
        self.brightness_slider.valueChanged.connect(
            self.refresh_background_preview
        )
        self.contrast_slider.valueChanged.connect(
            self.refresh_background_preview
        )

        self.sequence_button.clicked.connect(
            self.open_sequence_editor
        )

    # =========================================================
    # SLIDE MANAGEMENT
    # =========================================================

    def add_slide(self):
        slide = Slide()

        self.slides.append(slide)

        self.slide_list.addItem(
            self._build_slide_label(len(self.slides) - 1, slide)
        )

        self.slide_list.setCurrentRow(
            len(self.slides) - 1
        )

    def delete_slide(self):
        index = self.slide_list.currentRow()

        if index < 0:
            return

        self.slides.pop(index)
        self.slide_list.takeItem(index)

        if self.slides:
            new_index = min(
                index,
                len(self.slides) - 1,
            )

            self.slide_list.setCurrentRow(
                new_index
            )

        else:
            self.current_slide = None

    def select_slide(self, index):
        if index < 0 or index >= len(self.slides):
            self.current_slide = None
            return

        self.current_slide = self.slides[index]

        self.load_slide_into_editor(
            self.current_slide
        )

    def _build_slide_label(self, index, slide):
        label = f"{index + 1} - {slide.type.value.title()}"

        tags = []

        if slide.use_custom_style:
            tags.append("Custom Text")

        if slide.use_custom_background:
            tags.append("Custom BG")

        if tags:
            label += " [" + ", ".join(tags) + "]"

        return label

    # =========================================================
    # LOAD SLIDE
    # =========================================================

    def load_slide_into_editor(self, slide):
        # Type
        type_index = self.type_combo.findData(
            slide.type
        )

        if type_index >= 0:
            self.type_combo.blockSignals(True)
            self.type_combo.setCurrentIndex(
                type_index
            )
            self.type_combo.blockSignals(False)

        # Text content is always per-slide, regardless of custom flags.
        self.primary_text.setPlainText(
            slide.primary.content
        )

        self.secondary_text.setPlainText(
            slide.secondary.content
        )

        # Style: pull from the slide's own style only if it opted out of
        # the shared default.
        self.custom_style_checkbox.blockSignals(True)
        self.custom_style_checkbox.setChecked(
            slide.use_custom_style
        )
        self.custom_style_checkbox.blockSignals(False)

        primary_style = (
            slide.primary.style
            if slide.use_custom_style
            else self.global_defaults.primary_style
        )

        secondary_style = (
            slide.secondary.style
            if slide.use_custom_style
            else self.global_defaults.secondary_style
        )

        self._load_primary_style(primary_style)
        self._load_secondary_style(secondary_style)

        # Background: same idea.
        self.custom_background_checkbox.blockSignals(True)
        self.custom_background_checkbox.setChecked(
            slide.use_custom_background
        )
        self.custom_background_checkbox.blockSignals(False)

        background = (
            slide.background
            if slide.use_custom_background
            else self.global_defaults.background
        )

        self._apply_background_to_widgets(background)

    def _load_primary_style(self, style: Style):
        self.primary_font.setCurrentFont(
            QFont(style.font_family)
        )

        self.primary_size.setValue(
            style.font_size
        )

        self.primary_bold.setChecked(
            style.bold
        )

        self.primary_italic.setChecked(
            style.italic
        )

        index = self.primary_case.findData(style.case)
        self.primary_case.setCurrentIndex(
            index if index >= 0 else 0
        )

        self._pending_primary_color = style.font_color

        self.primary_color_button.setStyleSheet(
            f"background-color: {style.font_color};"
        )

    def _load_secondary_style(self, style: Style):
        self.secondary_font.setCurrentFont(
            QFont(style.font_family)
        )

        self.secondary_size.setValue(
            style.font_size
        )

        self.secondary_bold.setChecked(
            style.bold
        )

        self.secondary_italic.setChecked(
            style.italic
        )

        index = self.secondary_case.findData(style.case)
        self.secondary_case.setCurrentIndex(
            index if index >= 0 else 0
        )

        self._pending_secondary_color = style.font_color

        self.secondary_color_button.setStyleSheet(
            f"background-color: {style.font_color};"
        )

    def _apply_background_to_widgets(self, background: Background):
        self.brightness_slider.blockSignals(True)
        self.brightness_spin.blockSignals(True)
        self.brightness_slider.setValue(background.brightness)
        self.brightness_spin.setValue(background.brightness)
        self.brightness_slider.blockSignals(False)
        self.brightness_spin.blockSignals(False)

        self.contrast_slider.blockSignals(True)
        self.contrast_spin.blockSignals(True)
        self.contrast_slider.setValue(background.contrast)
        self.contrast_spin.setValue(background.contrast)
        self.contrast_slider.blockSignals(False)
        self.contrast_spin.blockSignals(False)

        self._pending_background_path = background.image_path

        self.background_path_label.setText(
            background.image_path.split("/")[-1]
            if background.image_path
            else "No image selected"
        )

        self.refresh_background_preview()

    # =========================================================
    # SAVE SLIDE
    # =========================================================

    def save_current_slide(self):
        if self.current_slide is None:
            return

        slide = self.current_slide

        # Type
        slide.type = self.type_combo.currentData()

        # Text content is always per-slide.
        slide.primary.content = (
            self.primary_text.toPlainText()
        )

        slide.secondary.content = (
            self.secondary_text.toPlainText()
        )

        # Style: write into the slide's own style if custom, otherwise
        # into the shared global default (which every non-custom slide
        # will then reflect).
        slide.use_custom_style = self.custom_style_checkbox.isChecked()

        primary_target = (
            slide.primary.style
            if slide.use_custom_style
            else self.global_defaults.primary_style
        )

        secondary_target = (
            slide.secondary.style
            if slide.use_custom_style
            else self.global_defaults.secondary_style
        )

        primary_target.font_family = (
            self.primary_font.currentFont().family()
        )
        primary_target.font_size = self.primary_size.value()
        primary_target.bold = self.primary_bold.isChecked()
        primary_target.italic = self.primary_italic.isChecked()
        primary_target.font_color = self._pending_primary_color
        primary_target.case = self.primary_case.currentData()

        secondary_target.font_family = (
            self.secondary_font.currentFont().family()
        )
        secondary_target.font_size = self.secondary_size.value()
        secondary_target.bold = self.secondary_bold.isChecked()
        secondary_target.italic = self.secondary_italic.isChecked()
        secondary_target.font_color = self._pending_secondary_color
        secondary_target.case = self.secondary_case.currentData()

        # Background: same pattern.
        slide.use_custom_background = (
            self.custom_background_checkbox.isChecked()
        )

        background_target = (
            slide.background
            if slide.use_custom_background
            else self.global_defaults.background
        )

        background_target.image_path = self._pending_background_path
        background_target.brightness = self.brightness_slider.value()
        background_target.contrast = self.contrast_slider.value()

        # Update sidebar label
        index = self.slide_list.currentRow()

        if index >= 0:
            self.slide_list.item(index).setText(
                self._build_slide_label(index, slide)
            )

    # =========================================================
    # SLIDE TYPE
    # =========================================================

    def update_slide_type(self):
        if self.current_slide is None:
            return

        self.current_slide.type = (
            self.type_combo.currentData()
        )

        index = self.slide_list.currentRow()

        if index >= 0:
            self.slide_list.item(index).setText(
                self._build_slide_label(index, self.current_slide)
            )

    # =========================================================
    # CUSTOM-FORMATTING TOGGLES
    # =========================================================

    def on_custom_style_toggled(self, checked):
        if self.current_slide is None:
            return

        slide = self.current_slide

        if checked:
            # Start this slide's own style as a copy of the current
            # default, then let the user diverge from there.
            slide.primary.style = replace(
                self.global_defaults.primary_style
            )
            slide.secondary.style = replace(
                self.global_defaults.secondary_style
            )

            self._load_primary_style(slide.primary.style)
            self._load_secondary_style(slide.secondary.style)

        else:
            # Drop back to showing (and, on Save, using) the shared default.
            self._load_primary_style(self.global_defaults.primary_style)
            self._load_secondary_style(self.global_defaults.secondary_style)

    def on_custom_background_toggled(self, checked):
        if self.current_slide is None:
            return

        slide = self.current_slide

        if checked:
            slide.background = replace(self.global_defaults.background)
            self._apply_background_to_widgets(slide.background)

        else:
            self._apply_background_to_widgets(self.global_defaults.background)

    # =========================================================
    # SAVE AS DEFAULT
    # =========================================================

    def set_style_as_default(self):
        self.global_defaults.primary_style = Style(
            font_family=self.primary_font.currentFont().family(),
            font_size=self.primary_size.value(),
            font_color=self._pending_primary_color,
            bold=self.primary_bold.isChecked(),
            italic=self.primary_italic.isChecked(),
            case=self.primary_case.currentData(),
        )

        self.global_defaults.secondary_style = Style(
            font_family=self.secondary_font.currentFont().family(),
            font_size=self.secondary_size.value(),
            font_color=self._pending_secondary_color,
            bold=self.secondary_bold.isChecked(),
            italic=self.secondary_italic.isChecked(),
            case=self.secondary_case.currentData(),
        )

    def set_background_as_default(self):
        self.global_defaults.background.image_path = (
            self._pending_background_path
        )
        self.global_defaults.background.brightness = (
            self.brightness_slider.value()
        )
        self.global_defaults.background.contrast = (
            self.contrast_slider.value()
        )

    # =========================================================
    # COLORS
    # =========================================================

    def choose_color(self, language):
        color = QColorDialog.getColor()

        if not color.isValid():
            return

        if language == "primary":
            self._pending_primary_color = color.name()

            self.primary_color_button.setStyleSheet(
                f"background-color: {color.name()};"
            )

        else:
            self._pending_secondary_color = color.name()

            self.secondary_color_button.setStyleSheet(
                f"background-color: {color.name()};"
            )

    # =========================================================
    # BACKGROUND
    # =========================================================

    def choose_background(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Choose Background Image",
            "",
            "Images (*.png *.jpg *.jpeg *.bmp *.webp)",
        )

        if not file_path:
            return

        self._pending_background_path = file_path

        self.background_path_label.setText(
            file_path.split("/")[-1]
        )

        self.refresh_background_preview()

    def _ensure_thumbnail_cache(self, path):
        if not path:
            self._cached_thumbnail_original = None
            self._cached_original_path = None
            return

        if (
            self._cached_original_path == path
            and self._cached_thumbnail_original is not None
        ):
            return

        original = QImage(path)

        if original.isNull():
            self._cached_thumbnail_original = None
            self._cached_original_path = None
            return

        # Downscale once per image selection so the live preview stays
        # fast even while dragging the slider.
        self._cached_thumbnail_original = original.scaled(
            self.background_preview.width(),
            self.background_preview.height(),
            Qt.KeepAspectRatioByExpanding,
            Qt.SmoothTransformation,
        )

        self._cached_original_path = path

    def refresh_background_preview(self):
        self._ensure_thumbnail_cache(self._pending_background_path)

        if self._cached_thumbnail_original is None:
            self.background_preview.setPixmap(QPixmap())
            self.background_preview.setText("No Image")
            return

        adjusted = apply_brightness_contrast(
            self._cached_thumbnail_original,
            self.brightness_slider.value(),
            self.contrast_slider.value(),
        )

        self.background_preview.setText("")
        self.background_preview.setPixmap(
            QPixmap.fromImage(adjusted)
        )

    # =========================================================
    # PREVIEW
    # =========================================================

    def preview_slide(self):
        if self.current_slide is None:
            return

        self.save_current_slide()

        self.preview_window = self.renderer.render(
            self.current_slide,
            self.global_defaults,
            1280,
            720,
        )

        self.preview_window.setWindowTitle(
            "Worship Presenter - Preview"
        )

        self.preview_window.show()

    def save_project_dialog(self):
        self.save_current_slide()

        if self.current_project_path:
            file_path = self.current_project_path
        else:
            file_path, _ = QFileDialog.getSaveFileName(
                self,
                "Save Worship Project",
                "",
                "Worship Project (*.json)",
            )

            if not file_path:
                return

        if not file_path.lower().endswith(".json"):
            file_path += ".json"

        try:
            save_project(
                file_path,
                self.slides,
                self.global_defaults,
                self.presentation,
            )

            self.current_project_path = file_path

            self.statusBar().showMessage(
                "Project saved successfully.",
                5000,
            )

        except Exception as error:
            from PySide6.QtWidgets import QMessageBox

            QMessageBox.critical(
                self,
                "Save Error",
                f"Could not save project:\n{error}",
            )

    def open_project_dialog(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Open Worship Project",
            "",
            "Worship Project (*.json)",
        )

        if not file_path:
            return

        try:
            slides, global_defaults, presentation = load_project(
                file_path
            )

            self.slides = slides
            self.global_defaults = global_defaults
            self.presentation = presentation
            self.current_project_path = file_path

            # Close any open sequence editor.
            if self.sequence_window is not None:
                self.sequence_window.close()
                self.sequence_window = None

            # Refresh the slide list.
            self.slide_list.clear()

            for index, slide in enumerate(self.slides):
                self.slide_list.addItem(
                    self._build_slide_label(index, slide)
                )

            # Select the first slide, if available.
            if self.slides:
                self.slide_list.setCurrentRow(0)
            else:
                self.current_slide = None
                self._load_primary_style(
                    self.global_defaults.primary_style
                )
                self._load_secondary_style(
                    self.global_defaults.secondary_style
                )
                self._apply_background_to_widgets(
                    self.global_defaults.background
                )

            self.statusBar().showMessage(
                "Project loaded successfully.",
                5000,
            )

        except Exception as error:
            from PySide6.QtWidgets import QMessageBox

            QMessageBox.critical(
                self,
                "Open Error",
                f"Could not open project:\n{error}",
            )

    def open_sequence_editor(self):

        self.save_current_slide()

        # If a sequence window is already open, close it first instead of
        # stacking a second one on top -- avoids confusing duplicate
        # windows editing the same underlying presentation.
        if self.sequence_window is not None:
            self.sequence_window.close()
            self.sequence_window = None

        self.sequence_window = SequenceEditor(
            self.slides,
            self.presentation,
            self.renderer,
            self.global_defaults,
            self,
        )

        self.sequence_window.show()
        self.sequence_window.raise_()
        self.sequence_window.activateWindow()

    def open_presenter_defaults(self):
        self.save_current_slide()

        dialog = PresenterDefaultsWindow(
            self.global_defaults,
            self.presentation,
            self
        )

        if dialog.exec():
            self._load_primary_style(
                self.global_defaults.primary_style
            )
            self._load_secondary_style(
                self.global_defaults.secondary_style
            )
            self._apply_background_to_widgets(
                self.global_defaults.background
            )

            self.statusBar().showMessage(
                "Presenter defaults updated.", 5000
            )

    def export_powerpoint(self):
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export PowerPoint",
            "",
            "PowerPoint Presentation (*.pptx)"
        )

        if not file_path:
            return

        if not file_path.endswith(".pptx"):
            file_path += ".pptx"

        self.create_powerpoint(file_path)

    # Pixel size used to rasterize backgrounds for export. Matches the
    # slide's 13.333x7.5in (16:9) aspect ratio so SlideRenderer.create_background
    # -- the same scale/crop/center/brightness/contrast pipeline used for the
    # on-screen preview -- can be reused as-is and dropped in full-bleed.
    _EXPORT_BG_PIXEL_SIZE = (1920, 1080)

    def _set_font_color(self, font, hex_color):
        color = hex_color.lstrip("#")

        font.color.rgb = RGBColor(
            int(color[0:2], 16),
            int(color[2:4], 16),
            int(color[4:6], 16),
        )

    def _add_slide_background(self, ppt, slide, background, temp_image_paths):
        # Reuse the renderer's own background pipeline (scale, center-crop,
        # brightness/contrast) so the exported slide matches the preview
        # exactly, including the plain black fallback when there's no image.
        width_px, height_px = self._EXPORT_BG_PIXEL_SIZE

        background_pixmap = self.renderer.create_background(
            background, width_px, height_px
        )

        temp_file = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
        temp_file.close()

        background_pixmap.save(temp_file.name, "PNG")
        temp_image_paths.append(temp_file.name)

        slide.shapes.add_picture(
            temp_file.name,
            0, 0,
            width=ppt.slide_width,
            height=ppt.slide_height,
        )

    def _add_mismatch_error_textbox(self, slide):
        # Mirrors the on-screen renderer's behavior when primary/secondary
        # line counts don't match, instead of silently mangling the slide.
        error_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(3.0),
            Inches(12.33), Inches(1.5),
        )

        paragraph = error_box.text_frame.paragraphs[0]
        paragraph.text = (
            "ERROR: Primary and Secondary language "
            "line counts do not match."
        )
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.size = Pt(28)
        paragraph.font.bold = True
        self._set_font_color(paragraph.font, "#FF0000")

    def _add_slide_text(
        self,
        slide,
        primary_lines,
        secondary_lines,
        primary_style,
        secondary_style,
    ):
        if not primary_lines and not secondary_lines:
            return

        textbox = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.0),
            Inches(12.33), Inches(5.5),
        )

        text_frame = textbox.text_frame
        text_frame.word_wrap = True

        first_paragraph = True
        last_index = len(primary_lines) - 1

        for index, (primary_line, secondary_line) in enumerate(
            zip(primary_lines, secondary_lines)
        ):
            for content, style in (
                (primary_line, primary_style),
                (secondary_line, secondary_style),
            ):
                if first_paragraph:
                    paragraph = text_frame.paragraphs[0]
                    first_paragraph = False
                else:
                    paragraph = text_frame.add_paragraph()

                paragraph.text = content
                paragraph.alignment = PP_ALIGN.CENTER

                font = paragraph.font
                font.name = style.font_family
                font.size = Pt(style.font_size)
                font.bold = style.bold
                font.italic = style.italic
                self._set_font_color(font, style.font_color)

            # Small spacer between line-pairs, but not after the last one --
            # mirrors the vertical stack of primary/secondary label pairs in
            # SlideRenderer.render.
            if index != last_index:
                spacer = text_frame.add_paragraph()
                spacer.font.size = Pt(12)

    def create_powerpoint(self, file_path):
        from PySide6.QtWidgets import QMessageBox

        temp_image_paths = []

        try:
            self.save_current_slide()

            ppt = PptxPresentation()
            ppt.slide_width = Inches(13.333)
            ppt.slide_height = Inches(7.5)

            slides = SequenceResolver(self.slides).resolve_presentation(
                self.presentation
            )

            for slide_data in slides:
                slide = ppt.slides.add_slide(
                    ppt.slide_layouts[6]
                )

                background = (
                    slide_data.background
                    if slide_data.use_custom_background
                    else self.global_defaults.background
                )

                primary_style = (
                    slide_data.primary.style
                    if slide_data.use_custom_style
                    else self.global_defaults.primary_style
                )

                secondary_style = (
                    slide_data.secondary.style
                    if slide_data.use_custom_style
                    else self.global_defaults.secondary_style
                )

                self._add_slide_background(
                    ppt, slide, background, temp_image_paths
                )

                # Apply the same per-line text-case transform as the live
                # renderer (upper/lower/sentence/title), which the old
                # export path never applied.
                primary_lines = [
                    apply_text_case(line, primary_style.case)
                    for line in slide_data.primary.content.splitlines()
                ]

                secondary_lines = [
                    apply_text_case(line, secondary_style.case)
                    for line in slide_data.secondary.content.splitlines()
                ]

                if len(primary_lines) != len(secondary_lines):
                    self._add_mismatch_error_textbox(slide)
                    continue

                self._add_slide_text(
                    slide,
                    primary_lines,
                    secondary_lines,
                    primary_style,
                    secondary_style,
                )

            ppt.save(file_path)

            QMessageBox.information(
                self,
                "Export Complete",
                "PowerPoint exported successfully!",
            )

        except Exception as error:
            QMessageBox.critical(
                self,
                "Export Error",
                f"Could not export PowerPoint:\n{error}",
            )

        finally:
            for path in temp_image_paths:
                try:
                    os.remove(path)
                except OSError:
                    pass


def main():
    app = QApplication(sys.argv)

    window = MainWindow()
    window.show()

    sys.exit(app.exec())


if __name__ == "__main__":
    main()